"""Fill the official BioHack 2026 proposal template with HeatSignal NYC content."""
import copy
import sys
from pathlib import Path

import docx
from docx.shared import Pt

TEMPLATE = Path(sys.argv[1])
OUT = Path(__file__).resolve().parent.parent / "submission" / "HeatSignal_NYC_BioHack2026_Proposal.docx"
SIZE = Pt(8.5)

TEAM = "[Team name]"
MEMBERS = "[Member names]"

SECTIONS = [
    # Specific research question
    [[("", "NYC issues a heat advisory only when the heat index hits 100°F, or 95°F two days running [2]. Does that rule warn on the days New Yorkers actually get sick "
       "from heat? Can a forecast of heat ER visits, tied to neighborhood vulnerability, warn earlier and send low-cost help where it is needed most?")]],
    # Background
    [[("", "Heat is NYC's deadliest weather hazard: about 500 premature deaths a summer. Black New Yorkers die of heat stress at twice the white rate; most heat-stress deaths occur in homes without AC [1]. "
       "Daily heat ER visits track excess deaths [4]. The 95 to 100°F range matches a study of 1997 to 2006 deaths that urged re-checking with new data [5], and early heat waves are about twice as deadly [6]. "
       "NWS HeatRisk is national and climate-based, not tuned to NYC outcomes [7]. "),
      ("We found: 56% of 7,619 heat ER visits (2017 to 2026) fell on days with no advisory.", " It missed 61 of 158 high-burden days. At 91 to 97°F, early-season days had 35% more visits.")]],
    # Proposed solution
    [[("HeatSignal NYC", " turns each morning's NWS forecast into expected heat ER visits (with an 80% range) for 3 days. It adds a "), ("Heat Watch", " tier (9.35+ expected visits, about 20 days a summer) "
       "under the unchanged advisory. A Watch triggers low-cost actions in the 81 neighborhoods with Heat Vulnerability Index 4 or 5: calls to older adults living alone, texts, hospital heads-ups "
       "and longer cooling-center hours, guided by a ranked call list. "), ("New: ", "a city trigger calibrated to NYC's own ER outcomes, stated in expected patients, with each action's cost matched to how certain the alert is.")]],
    # Data & methodology
    [[("Data: ", "NYC Health Dept. (public) [3]: daily heat ER visits from syndromic surveillance of every NYC ER, and max heat index at LaGuardia, May to Sept. 2017 to 2026 (1,520 days); HVI for 198 neighborhoods. "),
      ("Method: ", "Poisson regression on heat index (extra slope above 85/90/95°F), prior-day heat, hot streaks, early-season heat, weekend, and visits 2 days earlier. Leave-one-summer-out validation with cutoffs set "
       "on training years, ±2.5°F forecast error simulated, and a script that re-checks all 44 numbers we cite. "),
      ("Results: ", "mean forecast matches mean actual in every forecast range (e.g., 14.3 vs. 14.8). Watch + Advisory warned on "),
      ("131 of 158 bad days (83%, 95% CI 75 to 91%) vs. 97 (61%, 51 to 74%)", ". A warning was in place from day one in 46 vs. 30 of 62 heat episodes, and in 82% of bad days with forecast error.")]],
    # Prototype plan
    [[("Built: ", "an in-browser app with a 3-day planner (live NWS forecast), an out-of-sample replay of any summer (June 2024: Watch a day before the advisory), and a neighborhood map with call list. "),
      ("Next: ", "ZIP-level forecasts with Health Dept. data, adding overnight heat and ozone, and a one-summer South Bronx pilot measuring calls made and ER visits.")]],
    # Impact
    [[("", "Users: Health Dept., NYC Emergency Management, outreach groups, hospitals. Beneficiaries: South Bronx, Upper Manhattan and central Brooklyn residents (HVI-5 areas as low as 76% AC). "
       "Over 10 summers, about 1,000 more heat ER visits fell on warned days (about 100 a summer) for about 8 extra alert days a year. A false Watch costs some calls; a miss can cost a life. "
       "Any city among the 85% of US ERs reporting to CDC syndromic surveillance [8] could reuse it.")]],
    # Feasibility
    [[("", "Built this weekend from public data, no servers. Roles: [Member 1] data and model; [Member 2] validation; [Member 3] web app; [Member 4] playbook and pitch. "
       "Deploying needs no new data system: the city gets these counts daily. The model is refit each fall.")]],
    # Limitations
    [[("", "Counts are suspected illness and miss deaths at home and people who avoid ERs. One weather station; citywide outcomes only. A plain lower heat-index cutoff gets close (129 vs. 131), "
       "so most gain comes from warning at lower temperatures; the model adds patient estimates and precision (66% vs. 61%). It underpredicts the worst early heat waves "
       "(June 25, 2025: 31 expected, 141 actual). Watch goes to outreach teams, not sirens, to limit alert fatigue.")]],
]

REFS = [
    "[1] NYC DOHMH. 2026 Heat-Related Mortality Report. a816-dohbesp.nyc.gov/IndicatorPublic/data-features/heat-report  [2] National Weather Service New York. Heat Advisory criteria. weather.gov/okx/extremeheat",
    "[3] NYC DOHMH Environment & Health Data Portal data (heat-syndrome ED visits; Heat Vulnerability Index). github.com/nychealth/EHDP-data, accessed Oct 3, 2026.",
    "[4] Mathes RW et al. PLoS One. 2017;12(9):e0184364. doi:10.1371/journal.pone.0184364  [5] Metzger KB, Ito K, Matte TD. Environ Health Perspect. 2010;118(1):80-86. doi:10.1289/ehp.0900906",
    "[6] Anderson GB, Bell ML. Environ Health Perspect. 2011;119(2):210-218. doi:10.1289/ehp.1002313  [7] NOAA/CDC. NWS HeatRisk (2024), noaa.gov  [8] CDC. About NSSP, cdc.gov/nssp",
    "Live forecast: NWS API. Tools: Python (pandas, scikit-learn), JavaScript; Claude (Anthropic) AI for coding and drafting. Code and fact check (FACTS.md) included with the submission files.",
]


def set_runs(par, parts, size=SIZE):
    for r in list(par.runs):
        r._r.getparent().remove(r._r)
    for bold_txt, txt in parts:
        if bold_txt:
            r = par.add_run(bold_txt); r.bold = True; r.font.size = size
        if txt:
            r = par.add_run(txt); r.font.size = size


def main():
    d = docx.Document(TEMPLATE)
    title = d.paragraphs[0]
    for r in title.runs:
        if "[Insert Project Title]" in r.text:
            r.text = r.text.replace("[Insert Project Title]", "HeatSignal NYC")
    info = d.paragraphs[1]
    for r in info.runs:
        r.text = (r.text.replace("[Enter team name]", TEAM).replace("[Names of members]", MEMBERS)
                  .replace("[Stream name]", "Public Health & Epidemiology"))
    for tbl, paras in zip(d.tables, SECTIONS):
        cell = tbl.rows[0].cells[0]
        ps = cell.paragraphs
        for extra in ps[1:]:
            extra._p.getparent().remove(extra._p)
        set_runs(ps[0], paras[0])
        pf = ps[0].paragraph_format; pf.space_before = Pt(0); pf.space_after = Pt(0)
    for p in d.paragraphs[2:]:  # tighten the template's gaps above section headings
        p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(1)
    # references go in the empty paragraphs after the "Cite all references" heading
    tail = [p for p in d.paragraphs[-2:]]
    anchor = tail[0]
    set_runs(anchor, [("", REFS[0])], Pt(6.5))
    anchor.paragraph_format.space_before = Pt(0); anchor.paragraph_format.space_after = Pt(0)
    prev = anchor
    for ref in REFS[1:]:
        new = copy.deepcopy(anchor._p)
        prev._p.addnext(new)
        prev = docx.text.paragraph.Paragraph(new, anchor._parent)
        set_runs(prev, [("", ref)], Pt(6.5))
    tail[1]._p.getparent().remove(tail[1]._p)
    d.core_properties.author = "HeatSignal NYC team"
    d.core_properties.title = "HeatSignal NYC proposal"
    d.core_properties.last_modified_by = "HeatSignal NYC team"
    d.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
